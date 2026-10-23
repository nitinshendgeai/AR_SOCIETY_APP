"""Reading an asset register from an Excel (.xlsx) or CSV file.

Pure parsing, no database: `read_table` turns the file into a header row and data rows, `auto_map` matches its
headings to the asset fields (by the many names a register uses for them), and `parse_row` turns one row into the
values for an Asset or the reasons it can't be. InventoryService.import_assets does the checking against the society's
existing assets and the saving.

Every column in the file is read; the ones that match a field are imported and the rest are listed as ignored, so the
society can import its own sheet as it is. A column can be pointed at a different field (or ignored) by the caller.
"""
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from app.modules.inventory.models.inventory import AssetCategory, AssetStatus

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
HEADER_SEARCH_ROWS = 10      # a title above the headings is skipped: the best-matching of the first rows is the header


class ImportFileError(ValueError):
    """The file can't be read at all (wrong type, empty, too big, no recognisable columns)."""


# field -> (label, [heading names, written the way _key() writes them])
FIELDS: Dict[str, Tuple[str, List[str]]] = {
    "asset_code": ("Asset code", ["assetcode", "assetid", "assetno", "assetnumber", "code", "tagno", "tagnumber",
                                  "tag", "assettag"]),
    "name": ("Name", ["name", "assetname", "asset", "item", "itemname", "equipment", "equipmentname",
                      "nameofasset", "nameofequipment", "particulars"]),
    "asset_category": ("Category", ["category", "assetcategory", "type", "assettype", "kind", "group",
                                    "assetgroup", "assetclass", "class"]),
    "description": ("Description", ["description", "details", "remarks", "notes", "specification", "specifications",
                                    "specs", "comments"]),
    "location": ("Location", ["location", "place", "installedat", "installedin", "area", "site", "position",
                              "placement", "locationofasset"]),
    "status": ("Status", ["status", "assetstatus", "condition", "workingstatus"]),
    "purchase_date": ("Purchase date", ["purchasedate", "dateofpurchase", "purchasedon", "dop", "boughton",
                                        "installationdate", "dateofinstallation", "installeddate",
                                        "commissioningdate", "dateofcommissioning", "acquisitiondate"]),
    "purchase_cost": ("Purchase cost", ["purchasecost", "cost", "price", "amount", "value", "purchaseprice",
                                        "purchasevalue", "costinr", "costrs", "costinrs", "assetvalue", "rate"]),
    "vendor_name": ("Vendor", ["vendor", "vendorname", "supplier", "suppliername", "purchasedfrom", "dealer",
                               "boughtfrom", "vendorsupplier"]),
    "vendor_contact": ("Vendor contact", ["vendorcontact", "vendorphone", "vendormobile", "vendorno",
                                          "suppliercontact", "supplierphone", "suppliermobile", "contact",
                                          "contactno", "contactnumber", "phone", "mobile", "vendorcontactno"]),
    "invoice_number": ("Invoice number", ["invoice", "invoiceno", "invoicenumber", "billno", "billnumber",
                                          "invoiceref", "purchaseinvoice"]),
    "warranty_expiry": ("Warranty expiry", ["warranty", "warrantyexpiry", "warrantyuntil", "warrantyenddate",
                                            "warrantytill", "warrantyexpirydate", "warrantyvalidtill",
                                            "warrantyend", "warrantyexpireson", "warrantyupto"]),
    "expected_life_years": ("Expected life (years)", ["expectedlife", "expectedlifeyears", "life", "lifeyears",
                                                      "usefullife", "usefullifeyears", "lifespan",
                                                      "expectedlifeyr"]),
    "serial_number": ("Serial number", ["serial", "serialno", "serialnumber", "serialnum", "manufacturerserial",
                                        "serialnumberofasset"]),
    "model_number": ("Model number", ["model", "modelno", "modelnumber", "modelname", "modelnum"]),
    "service_interval_months": ("Service interval (months)", ["serviceinterval", "serviceintervalmonths",
                                                              "servicefrequency", "maintenanceinterval",
                                                              "servicecycle", "maintenancefrequency",
                                                              "serviceevery"]),
    "last_serviced_on": ("Last serviced on", ["lastserviced", "lastservicedon", "lastservicedate", "lastservice",
                                              "dateoflastservice", "lastmaintenance", "lastmaintenancedate"]),
    "next_service_due": ("Next service due", ["nextservice", "nextservicedue", "nextservicedate", "servicedue",
                                              "nextduedate", "nextmaintenance", "nextmaintenancedate",
                                              "servicedueon"]),
}
FIELD_LABELS = {k: v[0] for k, v in FIELDS.items()}
_BY_KEY = {key: field for field, (_, keys) in FIELDS.items() for key in keys}

MAX_LEN = {"asset_code": 30, "name": 255, "location": 255, "vendor_name": 255, "vendor_contact": 100,
           "invoice_number": 100, "serial_number": 100, "model_number": 100}
DATE_FIELDS = ("purchase_date", "warranty_expiry", "last_serviced_on", "next_service_due")


def _key(heading: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(heading).lower()) if heading is not None else ""


def _text(v: Any) -> str:
    """A cell as text: whole-number floats lose their '.0', dates read ISO, blanks are ''."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return re.sub(r"[\x00-\x08\x0b-\x1f]", "", str(v)).strip()


# ── Reading the file ──────────────────────────────────────────────────────────

def read_table(data: bytes, filename: str) -> Tuple[List[Any], List[Tuple[int, List[Any]]], str]:
    """(header cells, [(line number in the file, cells)], sheet name) from an .xlsx or .csv file. Blank rows are
    dropped; the line numbers are the file's own, so an error can be found in Excel."""
    if not data:
        raise ImportFileError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ImportFileError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. Split it and import in parts.")
    name = (filename or "").lower()
    if data[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        raw, sheet = _xlsx_rows(data)
    elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or name.endswith(".xls"):
        raise ImportFileError("Old Excel files (.xls) can't be read. Open it in Excel and Save As .xlsx or .csv.")
    else:
        raw, sheet = _csv_rows(data), "CSV"
    rows = [(n, r) for n, r in enumerate(raw, start=1) if any(_text(c) for c in r)]
    if not rows:
        raise ImportFileError("No data found in the file.")
    start = _header_row([r for _, r in rows])
    header, body = rows[start][1], rows[start + 1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"The file has {len(body)} rows; the limit is {MAX_ROWS} per import. Split it.")
    return header, body, sheet


def _csv_rows(data: bytes) -> List[List[Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _xlsx_rows(data: bytes) -> Tuple[List[List[Any]], str]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("That file could not be opened as an Excel (.xlsx) workbook.")
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(any(_text(c) for c in r) for r in rows):
                return rows, ws.title
    finally:
        wb.close()
    raise ImportFileError("No data found in the workbook.")


def _header_row(rows: List[List[Any]]) -> int:
    """The row of headings: of the first few rows, the one naming the most known fields (a title row above it, or a
    banner, is skipped). Falls back to the first row when none matches."""
    best, best_hits = 0, 0
    for i, r in enumerate(rows[:HEADER_SEARCH_ROWS]):
        hits = len({_BY_KEY[_key(c)] for c in r if _key(c) in _BY_KEY})
        if hits > best_hits:
            best, best_hits = i, hits
    return best


# ── Matching columns to fields ────────────────────────────────────────────────

def auto_map(header: List[Any]) -> List[Optional[str]]:
    """The field each column holds, or None when it isn't one we keep. A field is taken by the first column naming it;
    a later column with the same meaning is left unmapped."""
    taken, out = set(), []
    for cell in header:
        field = _BY_KEY.get(_key(cell))
        if field and field not in taken:
            taken.add(field)
            out.append(field)
        else:
            out.append(None)
    return out


def apply_override(header_map: List[Optional[str]], header: List[Any], override: Optional[Dict[str, Optional[str]]]
                   ) -> List[Optional[str]]:
    """The caller's choices (`{heading: field or null}`) win over the automatic match. Unknown fields are refused."""
    if not override:
        return header_map
    out = list(header_map)
    for i, cell in enumerate(header):
        label = _text(cell) or f"Column {_col_letter(i)}"
        if label in override:
            field = override[label]
            if field not in (None, "") and field not in FIELDS:
                raise ImportFileError(f"'{field}' is not a field that can be imported.")
            out[i] = field or None
    seen = set()
    for i, f in enumerate(out):
        if f:
            if f in seen:
                raise ImportFileError(f"Two columns are set to {FIELD_LABELS[f]}; each field can come from one column.")
            seen.add(f)
    return out


def _col_letter(i: int) -> str:
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def column_labels(header: List[Any]) -> List[str]:
    return [_text(c) or f"Column {_col_letter(i)}" for i, c in enumerate(header)]


# ── Values ────────────────────────────────────────────────────────────────────

_CATEGORY_WORDS = [
    ("lift", ["lift", "elevator", "escalator"]),
    ("cctv", ["cctv", "camera", "surveillance", "dvr", "nvr"]),
    ("pump", ["pump", "motor", "borewell", "booster"]),
    ("generator", ["generator", "dg", "genset", "diesel", "dg set"]),
    ("biometric", ["biometric", "attendance", "fingerprint"]),
    ("fire_safety", ["fire", "extinguisher", "sprinkler", "hydrant", "smoke"]),
    ("air_conditioner", ["air conditioner", "air conditioning", "ac", "a c", "split ac", "airconditioner"]),
    ("hvac", ["hvac", "ventilation", "exhaust", "blower"]),
    ("water_tank", ["water tank", "tank", "sump", "overhead tank"]),
    ("water_treatment", ["stp", "wtp", "etp", "ro", "water treatment", "sewage", "purifier", "softener"]),
    ("solar", ["solar", "inverter", "panel solar"]),
    ("gym_equipment", ["gym", "treadmill", "exercise", "fitness"]),
    ("intercom", ["intercom", "video door", "door phone"]),
    ("garden_equipment", ["garden", "lawn", "mower", "sprinkler garden"]),
    ("electrical", ["electrical", "electric", "transformer", "panel", "switchgear", "lighting", "light", "ups",
                    "battery", "meter"]),
    ("vehicle", ["vehicle", "car", "van", "truck", "tractor", "scooter", "bike", "golf cart"]),
    ("furniture", ["furniture", "chair", "table", "sofa", "bench", "cupboard", "almirah", "desk"]),
    ("it_equipment", ["it", "computer", "laptop", "printer", "server", "router", "network", "it equipment",
                      "pc", "desktop", "monitor"]),
    ("other", ["other", "others", "misc", "miscellaneous", "general"]),
]
_STATUS_WORDS = [
    ("under_maintenance", ["under maintenance", "maintenance", "repair", "under repair", "faulty", "not working",
                           "out of order", "breakdown", "service"]),
    ("retired", ["retired", "scrap", "scrapped", "obsolete", "decommissioned"]),
    ("disposed", ["disposed", "sold", "auctioned", "written off", "discarded"]),
    ("lost", ["lost", "missing", "stolen", "untraceable"]),
    ("active", ["active", "in use", "working", "good", "ok", "running", "operational", "functional", "available"]),
]


def _words(text: str) -> str:
    return " " + re.sub(r"[^a-z0-9]+", " ", text.lower()).strip() + " "


def _match(text: str, table, values) -> Optional[str]:
    w = _words(text)
    if not w.strip():
        return None
    for v in values:                                     # the value or its label, exactly
        if w.strip() == v.replace("_", " "):
            return v
    for value, words in table:                           # then a keyword anywhere in it: "Passenger Lift", "DG Set 125kVA"
        if any(f" {k} " in w for k in words):
            return value
    return None


def parse_category(text: str) -> Optional[AssetCategory]:
    found = _match(text, _CATEGORY_WORDS, [c.value for c in AssetCategory])
    return AssetCategory(found) if found else None


def parse_status(text: str) -> Optional[AssetStatus]:
    found = _match(text, _STATUS_WORDS, [s.value for s in AssetStatus])
    return AssetStatus(found) if found else None


_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y", "%d-%b-%Y",
                 "%d %b %Y", "%d-%b-%y", "%d %b %y", "%d %B %Y", "%d-%B-%Y", "%b %d, %Y", "%B %d, %Y", "%Y/%m/%d",
                 "%b-%Y", "%b %Y", "%B %Y", "%m/%Y", "%Y")


def parse_date(v: Any) -> Optional[date]:
    """Dates as Indian sheets write them (day first), ISO, '12-Mar-2021', an Excel date cell, or an Excel serial
    number. Returns None for something that isn't a date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = _text(v)
    if not text:
        return None
    if re.fullmatch(r"\d{4,6}(\.0+)?", text) and 20000 <= int(float(text)) <= 80000:     # Excel serial (1954 to 2118)
        return date(1899, 12, 30) + timedelta(days=int(float(text)))
    for fmt in _DATE_FORMATS:
        try:
            d = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        if fmt in ("%b-%Y", "%b %Y", "%B %Y", "%m/%Y", "%Y"):
            d = d.replace(day=1)
        if 1950 <= d.year <= 2100:
            return d
    return None


def parse_money(v: Any) -> Optional[Decimal]:
    text = _text(v)
    if not text:
        return None
    text = re.sub(r"(?i)\b(rs|inr)\b\.?|[₹,\s]", "", text)
    try:
        d = Decimal(text)
    except InvalidOperation:
        return None
    return d if d.is_finite() else None


_INTERVAL_WORDS = {"monthly": 1, "bimonthly": 2, "quarterly": 3, "halfyearly": 6, "half yearly": 6,
                   "biannual": 6, "semiannual": 6, "yearly": 12, "annual": 12, "annually": 12}


def parse_months(v: Any) -> Optional[int]:
    """'6', '6 months', 'quarterly', 'yearly', '1 year'."""
    text = _text(v).lower()
    if not text:
        return None
    if text in _INTERVAL_WORDS:
        return _INTERVAL_WORDS[text]
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(months?|mos?|years?|yrs?|y)?", text)
    if not m:
        return None
    n = float(m.group(1))
    return round(n * 12) if (m.group(2) or "").startswith(("y",)) else round(n)


def parse_years(v: Any) -> Optional[int]:
    text = _text(v).lower()
    if not text:
        return None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(years?|yrs?|y)?", text)
    return round(float(m.group(1))) if m else None


# ── One row ───────────────────────────────────────────────────────────────────

def parse_row(raw: List[Any], field_map: List[Optional[str]]) -> Tuple[Dict[str, Any], List[str], List[str]]:
    """(values for the Asset, errors, notes). Errors mean the row is not imported; notes are things filled in or
    guessed that the person should know (an unrecognised category became 'Other')."""
    cells: Dict[str, Any] = {}
    for i, field in enumerate(field_map):
        if field and i < len(raw):
            cells[field] = raw[i]
    values: Dict[str, Any] = {}
    errors: List[str] = []
    notes: List[str] = []

    for field, limit in MAX_LEN.items():
        text = _text(cells.get(field))
        if len(text) > limit:
            errors.append(f"{FIELD_LABELS[field]} is longer than {limit} characters")
        elif text:
            values[field] = text
    if "name" not in values and not any("Name" in e for e in errors):
        errors.append("Name is required")

    desc = _text(cells.get("description"))
    if desc:
        values["description"] = desc

    raw_cat = _text(cells.get("asset_category"))
    category = parse_category(raw_cat)
    if category is None:
        values["asset_category"] = AssetCategory.OTHER
        notes.append(f"Category '{raw_cat}' not recognised, imported as Other" if raw_cat
                     else "No category, imported as Other")
    else:
        values["asset_category"] = category

    raw_status = _text(cells.get("status"))
    status = parse_status(raw_status) if raw_status else AssetStatus.ACTIVE
    if status is None:
        status = AssetStatus.ACTIVE
        notes.append(f"Status '{raw_status}' not recognised, imported as Active")
    values["status"] = status

    for field in DATE_FIELDS:
        raw_v = cells.get(field)
        if _text(raw_v):
            d = parse_date(raw_v)
            if d is None:
                errors.append(f"{FIELD_LABELS[field]} '{_text(raw_v)}' is not a date (use dd/mm/yyyy)")
            else:
                values[field] = d

    if _text(cells.get("purchase_cost")):
        cost = parse_money(cells.get("purchase_cost"))
        if cost is None:
            errors.append(f"Purchase cost '{_text(cells.get('purchase_cost'))}' is not an amount")
        elif cost < 0:
            errors.append("Purchase cost can't be negative")
        elif cost >= Decimal("10000000000"):
            errors.append("Purchase cost is too large")
        else:
            values["purchase_cost"] = cost

    if _text(cells.get("service_interval_months")):
        months = parse_months(cells.get("service_interval_months"))
        if months is None or not 1 <= months <= 120:
            errors.append(f"Service interval '{_text(cells.get('service_interval_months'))}' must be 1 to 120 months")
        else:
            values["service_interval_months"] = months
    if _text(cells.get("expected_life_years")):
        years = parse_years(cells.get("expected_life_years"))
        if years is None or not 1 <= years <= 100:
            errors.append(f"Expected life '{_text(cells.get('expected_life_years'))}' must be 1 to 100 years")
        else:
            values["expected_life_years"] = years
    return values, errors, notes


# ── The template ──────────────────────────────────────────────────────────────

TEMPLATE_HEADINGS = ["Asset Code", "Name", "Category", "Location", "Status", "Purchase Date", "Purchase Cost",
                     "Vendor", "Vendor Contact", "Invoice Number", "Warranty Expiry", "Expected Life (Years)",
                     "Serial Number", "Model Number", "Service Interval (Months)", "Last Serviced On",
                     "Next Service Due", "Description"]
_TEMPLATE_SAMPLES = [
    ["", "Passenger lift - Wing A", "Lift", "Wing A lobby", "In use", "12/03/2021", "1250000", "Otis India",
     "9876543210", "INV-2021-118", "12/03/2026", "15", "SN-10045", "Gen2", "3", "10/01/2026", "", "Machine-room-less"],
    ["", "DG set 125 kVA", "Generator", "Basement", "In use", "05/04/2019", "800000", "Kirloskar", "", "", "",
     "10", "", "", "6", "", "", ""],
]


def template_csv() -> str:
    """A CSV with every importable column and two example rows. Only Name is required; the rest may be left blank
    (an empty Asset Code is numbered automatically)."""
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(TEMPLATE_HEADINGS)
    w.writerows(_TEMPLATE_SAMPLES)
    return out.getvalue()
